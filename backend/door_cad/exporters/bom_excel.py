"""Readable BOM workbook generated from canonical part rows."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from door_cad.models import ProjectGeometry
from door_cad.services.bom import BomRow, build_bom_rows


HEADERS = ["序号", "零件编号", "零件名称", "位置", "材料类型", "长度/mm", "展开宽/mm", "厚度/mm", "数量", "工艺备注"]


def build_bom_workbook(geometry: ProjectGeometry, rows: list[BomRow] | None = None) -> bytes:
    rows = rows if rows is not None else build_bom_rows(geometry)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "门框BOM"

    sheet.merge_cells("A1:J1")
    sheet["A1"] = "门框下料 BOM"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    metadata = [
        ("订单号", geometry.project.orderNo),
        ("项目名称", geometry.project.projectName),
        ("规则版本", geometry.ruleVersion),
        ("校验结果", geometry.validation.status),
    ]
    for index, (label, value) in enumerate(metadata, start=2):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=5)
        sheet.cell(index, 2, value)

    header_row = 7
    thin = Side(style="thin", color="808080")
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_index, row in enumerate(rows, start=header_row + 1):
        values = [
            row_index - header_row,
            row.partId,
            row.name,
            row.position,
            row.materialType,
            row.length,
            row.flatWidth,
            row.thickness,
            row.quantity,
            row.notes,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if column in (3, 10) else "center",
                vertical="center",
                wrap_text=column in (3, 10),
            )
        for column in (6, 7, 8):
            sheet.cell(row_index, column).number_format = "0.000"
        sheet.row_dimensions[row_index].height = 28 if row.notes else 22

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:J{header_row + len(rows)}"
    widths = {"A": 8, "B": 14, "C": 18, "D": 10, "E": 12, "F": 14, "G": 14, "H": 12, "I": 9, "J": 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
