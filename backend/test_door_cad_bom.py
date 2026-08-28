from io import BytesIO

from openpyxl import load_workbook

from door_cad.exporters.bom_excel import build_bom_workbook
from door_cad.models import FrameInput, ProjectMeta
from door_cad.services import build_bom_rows, calculate_frame_project
from test_door_cad_api import _client, _request


def _geometry():
    return calculate_frame_project(
        FrameInput(),
        ProjectMeta(orderNo="DD20260829002", projectName="BOM 测试"),
    )


def test_bom_rows_are_direct_projection_of_geometry():
    geometry = _geometry()
    rows = build_bom_rows(geometry)
    assert len(rows) == len(geometry.parts) == 8
    for row, part in zip(rows, geometry.parts):
        assert row.partId == part.partId
        assert row.length == part.length
        assert row.flatWidth == part.flatWidth
        assert row.thickness == part.thickness
        assert row.quantity == 1


def test_bom_workbook_reopens_with_readable_values():
    geometry = _geometry()
    workbook = load_workbook(BytesIO(build_bom_workbook(geometry)), data_only=False)
    sheet = workbook["门框BOM"]
    assert sheet["A1"].value == "门框下料 BOM"
    assert sheet["B2"].value == "DD20260829002"
    assert sheet.freeze_panes == "A8"
    assert sheet.auto_filter.ref == "A7:J15"
    assert sheet.max_row == 15
    assert [sheet.cell(row, 2).value for row in range(8, 16)] == [part.partId for part in geometry.parts]
    assert sheet.column_dimensions["J"].width >= 40


def test_bom_and_json_endpoints(tmp_path):
    client = _client(tmp_path)
    request = _request("导出测试")
    request["acknowledgeWarnings"] = True

    response = client.post("/api/door-cad/frame/export-bom", json=request)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["门框BOM"].max_row == 15

    response = client.post("/api/door-cad/frame/export-json", json=request)
    assert response.status_code == 200
    assert response.json()["schemaVersion"] == "1.0"
    assert len(response.json()["parts"]) == 8
