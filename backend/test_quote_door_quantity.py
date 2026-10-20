import copy
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from quote_database import QuoteDatabaseManager
from quote_models import QuoteCreate
from quote_excel import generate_excel


def payload(quantity=2):
    return {
        'customerName': 'Quantity test', 'quoteDate': '2026-10-18',
        'doorGroups': [{'groupName': 'Same specification', 'taskId': 'quantity-task',
                        'doorQuantity': quantity, 'items': [
                            {'productName': 'Door', 'unit': 'm2', 'width': 2580,
                             'height': 2890, 'unitPrice': 1480},
                            {'productName': 'Lock', 'unit': '套', 'quantity': 1, 'unitPrice': 580},
                        ]}],
    }


class QuoteDoorQuantityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = QuoteDatabaseManager(str(self.root / 'quotes.json'), str(self.root / 'backups'))

    def tearDown(self):
        self.temp.cleanup()

    def test_count_survives_save_reload_update_and_history_summary(self):
        created = self.db.create(payload())
        self.assertEqual(created['doorGroups'][0]['doorQuantity'], 2)
        self.assertEqual(self.db.get_by_id(created['id'])['doorGroups'][0]['doorQuantity'], 2)
        updated = self.db.update(created['id'], payload(3))
        self.assertEqual(self.db.get_all()[0]['doorCount'], 3)
        self.assertEqual(updated['doorGroups'][0]['items'][1]['quantity'], 1)
        self.assertEqual(self.db.find_groups_by_task('quantity-task')[0]['group']['doorQuantity'], 3)

    def test_legacy_groups_and_flat_items_default_to_one(self):
        data = payload()
        del data['doorGroups'][0]['doorQuantity']
        created = self.db.create(data)
        self.assertEqual(created['doorGroups'][0]['doorQuantity'], 1)
        raw = json.loads(Path(self.db.file_path).read_text(encoding='utf-8'))
        del raw[0]['doorGroups'][0]['doorQuantity']
        Path(self.db.file_path).write_text(json.dumps(raw), encoding='utf-8')
        self.assertEqual(self.db.get_by_id(created['id'])['doorGroups'][0]['doorQuantity'], 1)
        flat = self.db.create({'customerName': 'Legacy', 'quoteDate': '2026-10-18',
                               'items': data['doorGroups'][0]['items']})
        self.assertEqual(flat['doorGroups'][0]['doorQuantity'], 1)

    def test_request_and_storage_reject_invalid_counts(self):
        for value in (0, -1, 1.5, None, True, '2'):
            with self.subTest(value=value):
                with self.assertRaises(ValidationError):
                    QuoteCreate.model_validate(payload(value))
                with self.assertRaises(ValueError):
                    self.db.create(payload(value))

    def test_mixed_group_counts_sum_in_history(self):
        data = payload(2)
        data['doorGroups'].append(copy.deepcopy(data['doorGroups'][0]))
        data['doorGroups'][1]['doorQuantity'] = 3
        self.db.create(data)
        self.assertEqual(self.db.get_all()[0]['doorCount'], 5)

    def test_single_group_excel_subtotal_multiplies_once(self):
        quote = self.db.create(payload())
        filename = self.root / 'quote.xlsx'
        generate_excel(quote, str(filename))
        sheet = load_workbook(filename).worksheets[0]
        subtotals = [row for row in sheet.iter_rows() if '2 樘' in str(row[0].value)]
        self.assertEqual(len(subtotals), 1)
        self.assertIn('*2', subtotals[0][9].value)
        self.assertEqual(sheet['H10'].value, 1)

    def test_outer_area_screenshot_keeps_fraction_through_storage_and_excel(self):
        from quote_excel import _quote_quantity, _quote_amount
        data = payload(1)
        data['doorGroups'][0]['items'] = [{
            'productName': 'Outer area regression', 'width': 950, 'height': 2310,
            'unit': 'm2', 'unitPrice': 1480, 'quantity': None,
        }]
        saved = self.db.create(data)
        item = self.db.get_by_id(saved['id'])['doorGroups'][0]['items'][0]
        self.assertIsNone(item['quantity'])
        self.assertAlmostEqual(_quote_quantity(item), 2.1945)
        self.assertEqual(_quote_amount(item), 3248)
        filename = self.root / 'outer-area.xlsx'
        generate_excel(saved, str(filename))
        sheet = load_workbook(filename).worksheets[0]
        self.assertIn('D9*E9*0.000001', sheet['H9'].value)
        self.assertEqual(sheet['H9'].number_format, '0.####')
        self.assertIn('ROUND(H9*I9,0)', sheet['J9'].value)

    def test_order_quote_choice_keeps_per_door_amount_and_exposes_count(self):
        import sales_order_routes
        from unittest.mock import patch
        quote = self.db.create(payload())
        task = {'id': 'quantity-task', 'status': '已通过', 'params': {'dw': 2300, 'dh': 2750}}
        with patch.object(sales_order_routes, 'quote_db', self.db):
            candidate = sales_order_routes._candidate(task)
        self.assertEqual(candidate['quotes'][0]['door_quantity'], 2)
        self.assertEqual(candidate['quotes'][0]['amount'], 11615)
        self.assertEqual(candidate['quotes'][0]['items'][1]['quantity'], 1)

    def test_order_charges_round_per_door_before_extending_without_repricing_legacy(self):
        import sales_order_routes as routes
        from sales_order_models import SalesOrderChargeLineInput
        group = payload()['doorGroups'][0]
        group['items'] = [{'productName': 'Door', 'unit': 'm2', 'quantity': 2.52, 'unitPrice': 1680}]
        line = {'quantity': 2, 'quote_snapshot': {'group': group}}
        automatic = routes._charges_from_lines([line])
        self.assertEqual(automatic[0]['amount'], 8468)
        supplied = routes._charges_from_input([SalesOrderChargeLineInput(
            door_line_no=1, source_type='quote', product_name='Door', quantity=5.04,
            unit_price=1680)], 1, [line])
        self.assertEqual(supplied[0]['amount'], 8468)
        legacy = copy.deepcopy(line)
        del legacy['quote_snapshot']['group']['doorQuantity']
        self.assertEqual(routes._charges_from_lines([legacy])[0]['amount'], 8467)


if __name__ == '__main__':
    unittest.main()
