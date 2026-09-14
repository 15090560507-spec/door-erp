"""Previewed Excel/CSV imports for inventory master data."""

from __future__ import annotations

import csv
import io
import json
import sqlite3
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook

from inventory_database import InventoryDatabase, inventory_now
from inventory_service import InventoryService


ENTITY_CONFIG: Dict[str, Dict[str, Any]] = {
    "materials": {
        "label": "物料与商品",
        "fields": [
            ("code", "物料编码", True, ["编码", "商品编码", "物料编号", "itemcode"]),
            ("name", "物料名称", True, ["名称", "商品名称", "品名", "itemname"]),
            ("category", "分类", False, ["商品分类", "物料分类", "category"]),
            ("specification", "规格型号", False, ["规格", "型号", "spec", "specification"]),
            ("brand", "品牌", False, ["brand"]),
            ("unit", "基本单位", True, ["单位", "计量单位", "uom"]),
            ("material_type", "物料类型", True, ["类型", "商品类型", "itemtype"]),
            ("purchase_unit", "采购单位", False, ["采购计量单位"]),
            ("purchase_conversion", "采购换算率", False, ["单位换算", "换算率"]),
            ("default_supplier", "默认供应商", False, ["首选供应商"]),
            ("safety_stock", "安全库存", False, ["最低库存", "最小库存"]),
            ("standard_sale_price", "标准销售价", False, ["销售价", "零售价"]),
            ("reference_purchase_price", "参考采购价", False, ["采购价", "进价", "成本价"]),
            ("can_sell", "可销售", False, ["销售属性"]),
            ("can_purchase", "可采购", False, ["采购属性"]),
            ("manage_stock", "管理库存", False, ["库存管理"]),
            ("can_subcontract", "可外协", False, ["外协属性"]),
            ("remark", "备注", False, ["说明", "note"]),
        ],
    },
    "suppliers": {
        "label": "供应商",
        "fields": [
            ("code", "供应商编码", True, ["编码", "供应商编号", "vendorcode"]),
            ("name", "供应商名称", True, ["名称", "公司名称", "vendorname"]),
            ("short_name", "简称", False, ["供应商简称"]),
            ("contact_name", "联系人", False, ["联系人员"]),
            ("phone", "联系电话", False, ["电话", "手机"]),
            ("address", "地址", False, ["联系地址"]),
            ("invoice_title", "开票抬头", False, ["发票抬头"]),
            ("tax_no", "税号", False, ["纳税人识别号"]),
            ("default_tax_rate", "默认税率", False, ["税率", "税率%"]),
            ("settlement_method", "结算方式", False, ["付款方式"]),
            ("payment_days", "账期天数", False, ["账期"]),
            ("default_lead_days", "默认交期天数", False, ["交期天数", "交期"]),
            ("supply_category", "供货类别", False, ["供应分类"]),
            ("remark", "备注", False, ["说明", "note"]),
        ],
    },
    "supplier_items": {
        "label": "供应商产品",
        "fields": [
            ("supplier_code", "供应商编码", True, ["供应商编号", "vendorcode"]),
            ("material_code", "物料编码", True, ["商品编码", "物料编号", "itemcode"]),
            ("supplier_item_code", "供应商货号", False, ["供应商商品编码", "货号"]),
            ("supplier_item_name", "供应商商品名称", False, ["供应商品名"]),
            ("purchase_specification", "采购规格", False, ["供应商规格"]),
            ("purchase_unit", "采购单位", False, ["供应商单位"]),
            ("conversion_rate", "单位换算", False, ["换算率"]),
            ("tax_inclusive_price", "含税价", False, ["采购价", "含税单价", "进价"]),
            ("tax_rate", "税率", False, ["税率%"]),
            ("minimum_order_quantity", "最小采购量", False, ["起订量", "MOQ"]),
            ("lead_days", "交期天数", False, ["交期"]),
            ("is_preferred", "首选供应商", False, ["是否首选", "首选"]),
            ("remark", "备注", False, ["说明", "note"]),
        ],
    },
}


class MasterDataImportService:
    def __init__(self, db: InventoryDatabase):
        self.db = db
        self.inventory = InventoryService(db)

    def template(self, entity_type: str) -> Tuple[bytes, str]:
        config = self._config(entity_type)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = config["label"]
        sheet.append([field[1] for field in config["fields"]])
        if entity_type == "materials":
            sheet.append(["PLATE-08", "0.8mm不锈钢板", "板材", "1220x2440x0.8", "", "张", "原材料", "张", 1, "", 10, 0, 0, "否", "是", "是", "否", "示例行，可删除"])
        elif entity_type == "suppliers":
            sheet.append(["SUP-001", "示例供应商", "示例", "张经理", "13800000000", "", "", "", 13, "月结", 30, 7, "板材", "示例行，可删除"])
        else:
            sheet.append(["SUP-001", "PLATE-08", "TB-08", "0.8板材", "1220x2440", "张", 1, 680, 13, 5, 7, "是", "示例行，可删除"])
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(12, min(24, max(len(str(cell.value or "")) for cell in column) + 2))
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), f"{config['label']}_导入模板.xlsx"

    def preview(self, content: bytes, file_name: str, entity_type: str, created_by: str) -> Dict[str, Any]:
        config = self._config(entity_type)
        headers, rows = self._read_rows(content, file_name)
        if not rows:
            raise ValueError("导入文件没有可读取的数据行")
        mapping = self._auto_mapping(headers, config["fields"])
        now = inventory_now()
        with self.db.transaction() as conn:
            next_id = int(conn.execute("SELECT COALESCE(MAX(id), 0)+1 AS id FROM master_data_import_batches").fetchone()["id"])
            batch_no = f"DR{now[:10].replace('-', '')}{next_id:04d}"
            cursor = conn.execute(
                """INSERT INTO master_data_import_batches(
                       batch_no, entity_type, file_name, headers_json, rows_json, mapping_json,
                       total_rows, created_by, created_at, updated_at
                   ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (batch_no, entity_type, file_name, self._json(headers), self._json(rows), self._json(mapping), len(rows), created_by, now, now),
            )
            batch_id = int(cursor.lastrowid)
        return self.validate(batch_id, mapping, "skip")

    def validate(self, batch_id: int, mapping: Dict[str, str], duplicate_strategy: str) -> Dict[str, Any]:
        batch = self._batch(batch_id)
        config = self._config(batch["entity_type"])
        if duplicate_strategy not in {"skip", "update"}:
            raise ValueError("重复数据处理方式无效")
        headers = json.loads(batch["headers_json"])
        rows = json.loads(batch["rows_json"])
        allowed = {field[0] for field in config["fields"]}
        clean_mapping = {key: value for key, value in mapping.items() if key in allowed and value in headers}
        results: List[Dict[str, Any]] = []
        seen: set[str] = set()
        valid_rows = 0
        error_rows = 0
        for index, source in enumerate(rows, start=2):
            values = {key: source.get(column, "") for key, column in clean_mapping.items()}
            errors, warnings = self._validate_values(batch["entity_type"], values)
            key = self._key(batch["entity_type"], values)
            if key and key in seen:
                errors.append("文件内主键重复")
            if key:
                seen.add(key)
            if not errors and self._existing(batch["entity_type"], values):
                warnings.append("系统中已存在，将跳过" if duplicate_strategy == "skip" else "系统中已存在，将更新")
            if errors:
                error_rows += 1
            else:
                valid_rows += 1
            results.append({"row_number": index, "values": values, "errors": errors, "warnings": warnings})
        now = inventory_now()
        with self.db.transaction() as conn:
            conn.execute(
                """UPDATE master_data_import_batches SET mapping_json=?, duplicate_strategy=?,
                       valid_rows=?, error_rows=?, errors_json=?, updated_at=? WHERE id=?""",
                (self._json(clean_mapping), duplicate_strategy, valid_rows, error_rows, self._json([row for row in results if row["errors"]]), now, batch_id),
            )
        return {"batch": self.public_batch(batch_id), "fields": self.fields(batch["entity_type"]), "headers": headers, "mapping": clean_mapping, "rows": results[:50]}

    def execute(self, batch_id: int, mapping: Dict[str, str], duplicate_strategy: str, operator_uid: str) -> Dict[str, Any]:
        validation = self.validate(batch_id, mapping, duplicate_strategy)
        batch = self._batch(batch_id)
        if batch["status"] not in {"待执行", "预检失败"}:
            raise RuntimeError("该导入批次已经执行，不能重复导入")
        if validation["batch"]["error_rows"]:
            raise ValueError(f"预检仍有 {validation['batch']['error_rows']} 行错误，请修正文件或字段映射后再导入")
        rows = json.loads(batch["rows_json"])
        clean_mapping = validation["mapping"]
        imported = updated = skipped = 0
        failures: List[Dict[str, Any]] = []
        for index, source in enumerate(rows, start=2):
            values = {key: source.get(column, "") for key, column in clean_mapping.items()}
            try:
                existing = self._existing(batch["entity_type"], values)
                if existing and duplicate_strategy == "skip":
                    skipped += 1
                    continue
                target = self._save(batch["entity_type"], values, existing, operator_uid)
                action = "updated" if existing else "created"
                imported += int(not existing)
                updated += int(bool(existing))
                with self.db.transaction() as conn:
                    conn.execute(
                        """INSERT OR REPLACE INTO master_data_import_records(
                               batch_id, entity_type, target_id, target_key, action, created_at
                           ) VALUES (?, ?, ?, ?, ?, ?)""",
                        (batch_id, batch["entity_type"], target["id"], self._key(batch["entity_type"], values), action, inventory_now()),
                    )
            except Exception as exc:
                failures.append({"row_number": index, "message": str(exc)})
        status = "已完成" if not failures else "部分完成"
        now = inventory_now()
        with self.db.transaction() as conn:
            conn.execute(
                """UPDATE master_data_import_batches SET status=?, imported_rows=?, updated_rows=?,
                       skipped_rows=?, error_rows=?, errors_json=?, executed_at=?, updated_at=? WHERE id=?""",
                (status, imported, updated, skipped, len(failures), self._json(failures), now, now, batch_id),
            )
        return {"batch": self.public_batch(batch_id), "message": f"导入完成：新增 {imported}，更新 {updated}，跳过 {skipped}，失败 {len(failures)}"}

    def rollback(self, batch_id: int) -> Dict[str, Any]:
        batch = self._batch(batch_id)
        if batch["status"] not in {"已完成", "部分完成"}:
            raise RuntimeError("当前批次不能回滚")
        records = self.db.fetch_all(
            "SELECT * FROM master_data_import_records WHERE batch_id=? AND action='created' ORDER BY id DESC",
            (batch_id,),
        )
        deleted = 0
        blocked: List[Dict[str, Any]] = []
        tables = {"materials": "inventory_materials", "suppliers": "inventory_suppliers", "supplier_items": "inventory_supplier_items"}
        for record in records:
            try:
                with self.db.transaction() as conn:
                    conn.execute(f"DELETE FROM {tables[record['entity_type']]} WHERE id=?", (record["target_id"],))
                    conn.execute("DELETE FROM master_data_import_records WHERE id=?", (record["id"],))
                deleted += 1
            except sqlite3.IntegrityError:
                blocked.append({"target_key": record["target_key"], "message": "资料已被业务单据引用，未删除"})
        now = inventory_now()
        status = "已回滚" if not blocked else "部分回滚"
        with self.db.transaction() as conn:
            conn.execute(
                "UPDATE master_data_import_batches SET status=?, errors_json=?, rolled_back_at=?, updated_at=? WHERE id=?",
                (status, self._json(blocked), now, now, batch_id),
            )
        return {"batch": self.public_batch(batch_id), "message": f"回滚完成：删除 {deleted} 条，保留 {len(blocked)} 条已引用资料"}

    def list_batches(self, limit: int = 30) -> List[Dict[str, Any]]:
        rows = self.db.fetch_all("SELECT * FROM master_data_import_batches ORDER BY id DESC LIMIT ?", (limit,))
        return [self._public(row) for row in rows]

    def public_batch(self, batch_id: int) -> Dict[str, Any]:
        return self._public(self._batch(batch_id))

    @staticmethod
    def fields(entity_type: str) -> List[Dict[str, Any]]:
        config = ENTITY_CONFIG[entity_type]
        return [{"key": key, "label": label, "required": required} for key, label, required, _ in config["fields"]]

    def _save(self, entity_type: str, values: Dict[str, Any], existing: Dict[str, Any] | None, operator_uid: str) -> Dict[str, Any]:
        if entity_type == "materials":
            payload = self._material_payload(values)
            return self.inventory.update_material(existing["id"], payload) if existing else self.inventory.create_material(**{key: value for key, value in payload.items() if key != "is_active"})
        if entity_type == "suppliers":
            payload = self._supplier_payload(values)
            return self.inventory.update_supplier(existing["id"], payload) if existing else self.inventory.create_supplier(payload)
        supplier = self.db.fetch_one("SELECT id FROM inventory_suppliers WHERE code=?", (self._text(values.get("supplier_code")),))
        material = self.db.fetch_one("SELECT id FROM inventory_materials WHERE code=?", (self._text(values.get("material_code")),))
        payload = {
            "supplier_id": supplier["id"], "material_id": material["id"],
            "supplier_item_code": self._text(values.get("supplier_item_code")),
            "supplier_item_name": self._text(values.get("supplier_item_name")),
            "purchase_specification": self._text(values.get("purchase_specification")),
            "purchase_unit": self._text(values.get("purchase_unit")),
            "conversion_rate": self._number(values.get("conversion_rate"), 1),
            "tax_inclusive_price": self._number(values.get("tax_inclusive_price"), 0),
            "tax_rate": self._number(values.get("tax_rate"), 0),
            "minimum_order_quantity": self._number(values.get("minimum_order_quantity"), 0),
            "lead_days": int(self._number(values.get("lead_days"), 0)),
            "is_preferred": self._boolean(values.get("is_preferred"), False),
            "remark": self._text(values.get("remark")), "is_active": True,
        }
        return self.inventory.save_supplier_item(payload, operator_uid, existing["id"] if existing else None)

    def _material_payload(self, values: Dict[str, Any]) -> Dict[str, Any]:
        safety = self._number(values.get("safety_stock"), 0)
        return {
            "code": self._text(values.get("code")), "name": self._text(values.get("name")),
            "category": self._text(values.get("category")), "specification": self._text(values.get("specification")),
            "brand": self._text(values.get("brand")), "unit": self._text(values.get("unit")),
            "material_type": self._text(values.get("material_type")), "purchase_unit": self._text(values.get("purchase_unit")),
            "purchase_conversion": self._number(values.get("purchase_conversion"), 1),
            "default_supplier": self._text(values.get("default_supplier")), "minimum_stock": safety, "safety_stock": safety,
            "standard_sale_price": self._number(values.get("standard_sale_price"), 0),
            "reference_purchase_price": self._number(values.get("reference_purchase_price"), 0),
            "can_sell": self._boolean(values.get("can_sell"), False), "can_purchase": self._boolean(values.get("can_purchase"), True),
            "manage_stock": self._boolean(values.get("manage_stock"), True), "can_subcontract": self._boolean(values.get("can_subcontract"), False),
            "default_warehouse_id": None, "default_location_id": None, "remark": self._text(values.get("remark")), "is_active": True,
        }

    def _supplier_payload(self, values: Dict[str, Any]) -> Dict[str, Any]:
        return {
            "code": self._text(values.get("code")), "name": self._text(values.get("name")),
            "short_name": self._text(values.get("short_name")), "contact_name": self._text(values.get("contact_name")),
            "phone": self._text(values.get("phone")), "address": self._text(values.get("address")),
            "invoice_title": self._text(values.get("invoice_title")), "tax_no": self._text(values.get("tax_no")),
            "default_tax_rate": self._number(values.get("default_tax_rate"), 0),
            "settlement_method": self._text(values.get("settlement_method")),
            "payment_days": int(self._number(values.get("payment_days"), 0)),
            "default_lead_days": int(self._number(values.get("default_lead_days"), 0)),
            "supply_category": self._text(values.get("supply_category")), "remark": self._text(values.get("remark")), "is_active": True,
        }

    def _validate_values(self, entity_type: str, values: Dict[str, Any]) -> Tuple[List[str], List[str]]:
        errors: List[str] = []
        warnings: List[str] = []
        for key, label, required, _ in self._config(entity_type)["fields"]:
            if required and not self._text(values.get(key)):
                errors.append(f"{label}不能为空")
        numeric = {
            "materials": ["purchase_conversion", "safety_stock", "standard_sale_price", "reference_purchase_price"],
            "suppliers": ["default_tax_rate", "payment_days", "default_lead_days"],
            "supplier_items": ["conversion_rate", "tax_inclusive_price", "tax_rate", "minimum_order_quantity", "lead_days"],
        }[entity_type]
        for key in numeric:
            value = values.get(key)
            if self._text(value):
                try:
                    if float(value) < 0:
                        errors.append(f"{key}不能小于零")
                except (TypeError, ValueError):
                    errors.append(f"{key}必须是数字")
        if entity_type == "supplier_items":
            if self._text(values.get("supplier_code")) and not self.db.fetch_one("SELECT id FROM inventory_suppliers WHERE code=? AND is_active=1", (self._text(values.get("supplier_code")),)):
                errors.append("供应商编码不存在或已停用")
            if self._text(values.get("material_code")) and not self.db.fetch_one("SELECT id FROM inventory_materials WHERE code=? AND is_active=1", (self._text(values.get("material_code")),)):
                errors.append("物料编码不存在或已停用")
        return errors, warnings

    def _existing(self, entity_type: str, values: Dict[str, Any]) -> Dict[str, Any] | None:
        if entity_type == "materials":
            return self.db.fetch_one("SELECT * FROM inventory_materials WHERE code=?", (self._text(values.get("code")),))
        if entity_type == "suppliers":
            return self.db.fetch_one("SELECT * FROM inventory_suppliers WHERE code=?", (self._text(values.get("code")),))
        return self.db.fetch_one(
            """SELECT si.* FROM inventory_supplier_items si
               JOIN inventory_suppliers s ON s.id=si.supplier_id
               JOIN inventory_materials m ON m.id=si.material_id
               WHERE s.code=? AND m.code=?""",
            (self._text(values.get("supplier_code")), self._text(values.get("material_code"))),
        )

    @staticmethod
    def _key(entity_type: str, values: Dict[str, Any]) -> str:
        if entity_type in {"materials", "suppliers"}:
            return MasterDataImportService._text(values.get("code"))
        return f"{MasterDataImportService._text(values.get('supplier_code'))}|{MasterDataImportService._text(values.get('material_code'))}"

    @staticmethod
    def _auto_mapping(headers: List[str], fields: List[Tuple[str, str, bool, List[str]]]) -> Dict[str, str]:
        normalized = {MasterDataImportService._normalize(header): header for header in headers}
        mapping: Dict[str, str] = {}
        for key, label, _, aliases in fields:
            for candidate in [label, key, *aliases]:
                if MasterDataImportService._normalize(candidate) in normalized:
                    mapping[key] = normalized[MasterDataImportService._normalize(candidate)]
                    break
        return mapping

    @staticmethod
    def _read_rows(content: bytes, file_name: str) -> Tuple[List[str], List[Dict[str, Any]]]:
        suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
        matrix: List[List[Any]] = []
        if suffix == "csv":
            text = None
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ValueError("CSV 编码无法识别，请另存为 UTF-8 或 Excel 文件")
            matrix = [list(row) for row in csv.reader(io.StringIO(text))]
        elif suffix in {"xlsx", "xlsm"}:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            matrix = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        else:
            raise ValueError("仅支持 .xlsx、.xlsm 或 .csv 文件")
        if not matrix:
            raise ValueError("导入文件为空")
        headers = [str(value or "").strip() for value in matrix[0]]
        if not any(headers):
            raise ValueError("导入文件第一行必须是字段名称")
        named_headers = [header for header in headers if header]
        if len(named_headers) != len(set(named_headers)):
            raise ValueError("导入文件存在重复字段名称，请先修改表头")
        non_empty_rows = [values for values in matrix[1:] if any(value not in (None, "") for value in values)]
        if len(non_empty_rows) > 2000:
            raise ValueError("单次最多导入 2000 行，请拆分文件后重试")
        rows: List[Dict[str, Any]] = []
        for values in non_empty_rows:
            row = {header: MasterDataImportService._cell(values[index] if index < len(values) else "") for index, header in enumerate(headers) if header}
            rows.append(row)
        return headers, rows

    def _batch(self, batch_id: int) -> Dict[str, Any]:
        batch = self.db.fetch_one("SELECT * FROM master_data_import_batches WHERE id=?", (batch_id,))
        if not batch:
            raise LookupError("导入批次不存在")
        return batch

    @staticmethod
    def _public(batch: Dict[str, Any]) -> Dict[str, Any]:
        hidden = {"rows_json", "headers_json", "mapping_json", "errors_json"}
        result = {key: value for key, value in batch.items() if key not in hidden}
        result["errors"] = json.loads(batch.get("errors_json") or "[]")
        return result

    @staticmethod
    def _config(entity_type: str) -> Dict[str, Any]:
        if entity_type not in ENTITY_CONFIG:
            raise ValueError("不支持的基础资料类型")
        return ENTITY_CONFIG[entity_type]

    @staticmethod
    def _normalize(value: Any) -> str:
        return "".join(str(value or "").strip().lower().replace("（", "(").replace("）", ")").split())

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _number(value: Any, default: float) -> float:
        text = MasterDataImportService._text(value).replace(",", "").replace("%", "")
        return float(text) if text else default

    @staticmethod
    def _boolean(value: Any, default: bool) -> bool:
        text = MasterDataImportService._normalize(value)
        if not text:
            return default
        return text in {"1", "true", "yes", "y", "是", "启用", "首选"}

    @staticmethod
    def _cell(value: Any) -> Any:
        if value is None:
            return ""
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return value

    @staticmethod
    def _json(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
