"""Excel and print-document builders for production fulfillment records."""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from html import escape
from typing import Any, Dict, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from production_database import ProductionDatabase, json_loads, production_now


@dataclass
class ProductionDocument:
    title: str
    document_no: str
    metadata: List[Tuple[str, str]]
    headers: List[str]
    rows: List[List[Any]]
    orientation: str = "landscape"


def _order(db: ProductionDatabase, order_id: int) -> Dict[str, Any]:
    order = db.get_order(order_id)
    if not order:
        raise LookupError("生产订单不存在")
    return order


def build_order_document(
    db: ProductionDatabase, order_id: int, document_type: str
) -> ProductionDocument:
    order = _order(db, order_id)
    metadata = [
        ("生产单号", order["order_no"]),
        ("订货单位", order["customer"]),
        ("项目", order.get("project") or ""),
        ("要求交期", order.get("due_date") or ""),
    ]
    if document_type == "bom":
        state = db.fetch_one(
            "SELECT status FROM production_bom_status WHERE order_id=?", (order_id,)
        ) or {"status": "草稿"}
        rows = db.fetch_all(
            "SELECT * FROM production_bom_items WHERE order_id=? ORDER BY id", (order_id,)
        )
        return ProductionDocument(
            title="生产BOM",
            document_no=order["order_no"],
            metadata=metadata + [("BOM状态", state["status"])],
            headers=["序号", "分类", "物料名称", "规格", "材质", "厚度", "数量", "单位", "来源", "备注"],
            rows=[
                [index, row["category"], row["name"], row["specification"], row["material"],
                 row["thickness"], row["quantity"], row["unit"], row["supply_type"], row["remark"]]
                for index, row in enumerate(rows, 1)
            ],
        )
    if document_type == "cutting":
        sheet = db.fetch_one(
            "SELECT * FROM production_cutting_sheets WHERE order_id=?", (order_id,)
        )
        if not sheet:
            raise LookupError("综合下料单不存在")
        rows = db.fetch_all(
            "SELECT * FROM production_cutting_items WHERE sheet_id=? ORDER BY id", (sheet["id"],)
        )
        return ProductionDocument(
            title="综合下料单",
            document_no=order["order_no"],
            metadata=metadata + [("下料状态", sheet["status"])],
            headers=["序号", "材料", "规格", "理论数量", "实际数量", "单位", "下料人", "完成", "备注"],
            rows=[
                [index, row["name"], row["specification"], row["quantity"],
                 row["actual_quantity"], row["unit"], row["cutter"],
                 "是" if row["completed"] else "否", row["remark"]]
                for index, row in enumerate(rows, 1)
            ],
        )
    if document_type == "quality":
        rows = db.fetch_all(
            "SELECT * FROM production_quality_inspections WHERE order_id=? ORDER BY id DESC",
            (order_id,),
        )
        return ProductionDocument(
            title="成品质检单",
            document_no=order["order_no"],
            metadata=metadata,
            headers=["序号", "结果", "质检人", "退回工序", "备注", "质检时间"],
            rows=[
                [index, row["result"], row["inspector"], row["return_operation"],
                 row["remark"], row["created_at"]]
                for index, row in enumerate(rows, 1)
            ],
            orientation="portrait",
        )
    raise LookupError("不支持的生产单据类型")


def build_purchase_document(db: ProductionDatabase, purchase_id: int) -> ProductionDocument:
    purchase = db.fetch_one(
        "SELECT * FROM production_purchase_orders WHERE id=?", (purchase_id,)
    )
    if not purchase:
        raise LookupError("采购单不存在")
    rows = db.fetch_all(
        "SELECT * FROM production_purchase_items WHERE purchase_id=? ORDER BY id", (purchase_id,)
    )
    return ProductionDocument(
        title="采购单",
        document_no=purchase["purchase_no"],
        metadata=[
            ("采购单号", purchase["purchase_no"]),
            ("供应商", purchase["supplier"]),
            ("状态", purchase["status"]),
            ("预计到货", purchase["expected_date"]),
        ],
        headers=["序号", "生产单", "物料名称", "规格", "数量", "已到货", "单位", "单价", "金额", "备注"],
        rows=[
            [index, row["order_id"] or "", row["name"], row["specification"], row["quantity"],
             row["received_quantity"], row["unit"], row["unit_price"],
             round(float(row["quantity"]) * float(row["unit_price"]), 2), row["remark"]]
            for index, row in enumerate(rows, 1)
        ],
    )


def build_inventory_document(db: ProductionDatabase) -> ProductionDocument:
    rows = db.fetch_all(
        """
        SELECT t.*, m.code, m.name
        FROM production_inventory_transactions t
        LEFT JOIN production_materials m ON m.id=t.material_id
        ORDER BY t.id DESC
        """
    )
    return ProductionDocument(
        title="库存流水",
        document_no=production_now()[:10].replace("-", ""),
        metadata=[("导出时间", production_now()), ("流水条数", str(len(rows)))],
        headers=["序号", "物料编码", "物料名称", "生产订单", "类型", "数量", "单位", "仓位", "操作人", "时间", "备注"],
        rows=[
            [index, row.get("code") or "", row.get("name") or "成品", row.get("order_id") or "",
             row["transaction_type"], row["quantity"], row["unit"], row["warehouse_location"],
             row["operator_uid"], row["created_at"], row["remark"]]
            for index, row in enumerate(rows, 1)
        ],
    )


def build_shipment_document(db: ProductionDatabase, shipment_id: int) -> ProductionDocument:
    shipment = db.fetch_one("SELECT * FROM production_shipments WHERE id=?", (shipment_id,))
    if not shipment:
        raise LookupError("发货单不存在")
    rows = db.fetch_all(
        """
        SELECT o.order_no, o.customer, o.project, f.finished_no
        FROM production_shipment_items si
        JOIN production_orders o ON o.id=si.order_id
        JOIN production_finished_goods f ON f.id=si.finished_good_id
        WHERE si.shipment_id=? ORDER BY o.id
        """,
        (shipment_id,),
    )
    return ProductionDocument(
        title="成品发货单",
        document_no=shipment["shipment_no"],
        metadata=[
            ("发货单号", shipment["shipment_no"]),
            ("客户", shipment["customer"]),
            ("项目", shipment["project"]),
            ("状态", shipment["status"]),
            ("收货地址", shipment["address"]),
            ("联系人", f"{shipment['contact']} {shipment['phone']}".strip()),
            ("物流", shipment["logistics"]),
            ("运单号", shipment["tracking_no"]),
        ],
        headers=["序号", "生产单号", "成品编号", "订货单位", "项目", "数量", "单位"],
        rows=[
            [index, row["order_no"], row["finished_no"], row["customer"], row["project"], 1, "樘"]
            for index, row in enumerate(rows, 1)
        ],
        orientation="portrait",
    )


def render_xlsx(document: ProductionDocument) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = document.title[:31]
    column_count = max(1, len(document.headers))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title_cell = sheet.cell(1, 1, document.title)
    title_cell.font = Font(name="宋体", size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    row_index = 2
    for label, value in document.metadata:
        sheet.cell(row_index, 1, label).font = Font(name="宋体", bold=True)
        sheet.cell(row_index, 2, value)
        if column_count > 2:
            sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=column_count)
        row_index += 1

    header_row = row_index
    thin = Side(style="thin", color="808080")
    for column, label in enumerate(document.headers, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name="宋体", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for offset, values in enumerate(document.rows, 1):
        excel_row = header_row + offset
        max_lines = 1
        for column, value in enumerate(values, 1):
            cell = sheet.cell(excel_row, column, value)
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            max_lines = max(max_lines, math.ceil(len(str(value or "")) / 22))
        sheet.row_dimensions[excel_row].height = max(20, max_lines * 16)

    for column in range(1, column_count + 1):
        values: Sequence[Any] = [document.headers[column - 1]] + [
            row[column - 1] if column - 1 < len(row) else "" for row in document.rows
        ]
        longest = max((len(str(value or "")) for value in values), default=8)
        sheet.column_dimensions[get_column_letter(column)].width = min(36, max(10, longest + 2))

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{header_row + len(document.rows)}"
    sheet.print_area = f"A1:{get_column_letter(column_count)}{max(header_row + len(document.rows), header_row)}"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = document.orientation
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "第 &[Page] 页，共 &[Pages] 页"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_print_html(document: ProductionDocument) -> str:
    metadata = "".join(
        f"<div><b>{escape(str(label))}</b><span>{escape(str(value or ''))}</span></div>"
        for label, value in document.metadata
    )
    headers = "".join(f"<th>{escape(str(label))}</th>" for label in document.headers)
    rows = "".join(
        "<tr>" + "".join(f"<td>{escape(str(value or ''))}</td>" for value in row) + "</tr>"
        for row in document.rows
    )
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{escape(document.title)}</title>
<style>
@page {{ size: A4 {document.orientation}; margin: 10mm; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #111; font-family: "Noto Sans CJK SC", "Source Han Sans SC", "Microsoft YaHei", sans-serif; font-size: 9pt; }}
h1 {{ margin: 0 0 5mm; text-align: center; font-size: 16pt; }}
.meta {{ display: grid; grid-template-columns: 1fr 1fr; margin-bottom: 4mm; border-left: .5pt solid #555; border-top: .5pt solid #555; }}
.meta div {{ display: grid; grid-template-columns: 24mm 1fr; min-height: 8mm; border-right: .5pt solid #555; border-bottom: .5pt solid #555; }}
.meta b,.meta span {{ padding: 2mm; }}
table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
th,td {{ border: .5pt solid #555; padding: 1.8mm; overflow-wrap: anywhere; vertical-align: middle; }}
th {{ background: #d9eaf7; font-weight: 700; }}
.signatures {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 8mm; margin-top: 12mm; }}
button {{ position: fixed; top: 6mm; right: 6mm; padding: 2mm 5mm; }}
@media print {{ button {{ display: none; }} }}
</style></head><body><button onclick="window.print()">打印 / 保存PDF</button>
<h1>{escape(document.title)}</h1><section class="meta">{metadata}</section>
<table><thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table>
<section class="signatures"><span>制单：</span><span>审核：</span><span>日期：</span></section>
</body></html>"""
